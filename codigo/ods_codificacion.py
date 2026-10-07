"""
Codificación de las menciones ODS y fiabilidad entre codificadores.

Paso 1 · plantilla: junta las menciones web, los pasajes y los documentos en un Excel para codificar a mano.
  python codigo/ods_codificacion.py plantilla --web salidas/ods_menciones_web_FECHA.csv --docs salidas/ods_pasajes_FECHA.csv \
         --documentos salidas/ods_documentos_FECHA.csv --busquedas salidas/ods_busquedas_FECHA.csv
  → salidas/codificacion_ods_FECHA.xlsx con hojas:
     Menciones (Retórica/Sustantiva, dos codificadores y consenso) · Documentos (clasificación de cada
     ERD y cuenta pública, dos codificadores y consenso) · Mapa_17_ODS (nivel 0-3 de la ERD y 0/1/2 por ODS; filas
     cod. 1, cod. 2 y consenso) · Muestra_verificacion (5 GORE × términos para repetir la búsqueda web, semilla 2026)
  Las reglas de codificación están en «Reglas de medición» del README.
  La columna «Sugerencia automática» es solo una ayuda por palabras clave: NO reemplaza la codificación humana.

Paso 2 · fiabilidad: una vez que ambos codificadores llenaron el Excel (sin ver el trabajo del otro).
  python codigo/ods_codificacion.py kappa salidas/codificacion_ods_FECHA.xlsx
  → % de acuerdo y κ de Cohen (menciones, documentos y mapa) para revisar antes de cerrar la ronda.

Paso 3 · cerrar la ronda: python codigo/cerrar_ronda.py NOMBRE_RONDA --codificacion salidas/codificacion_ods_FECHA.xlsx
"""

import argparse
import random
import re

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from comun import SALIDAS, cargar_gores, hoy
from ronda import CLASES_DOC, kappa

VINCULO = re.compile(
    r"\b(eje|lineamiento|objetivo estrat[eé]gico|objetivos? espec[ií]fico|indicador|meta\s+\d|"
    r"financ|fondo conjunto|presupuesto|contribuye a|se relaciona con|alinead[oa]s? (con|a) (el|los) ods)",
    re.I,
)
F = Font(name="Arial", size=10)
FB = Font(name="Arial", size=10, bold=True)
FH = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HD = PatternFill("solid", fgColor="1C5CAB")
AMARILLO = PatternFill("solid", fgColor="FFF2CC")


def sugerencia(fila):
    """Heurística transparente: sustantiva si nombra ODS específicos junto a vocabulario de planificación/financiamiento."""
    ctx = str(fila.get("contexto", ""))
    especifico = bool(str(fila.get("ods_nombrados", "")).strip())
    return (
        "Sustantiva (revisar)"
        if VINCULO.search(ctx) and (especifico or "fondo conjunto" in ctx.lower())
        else "Retórica (revisar)"
    )


def encabezado(ws, fila, cols, anchos):
    """Escribe una fila de encabezado con formato y fija el ancho de las columnas."""
    for i, (c, w) in enumerate(zip(cols, anchos), 1):
        x = ws.cell(fila, i, c)
        x.font = FH
        x.fill = HD
        x.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[x.column_letter].width = w


def plantilla(a):
    """Comando «plantilla»: arma la planilla de codificación para dos codificadores."""
    web = pd.read_csv(a.web, encoding="utf-8-sig") if a.web else pd.DataFrame()
    docs = pd.read_csv(a.docs, encoding="utf-8-sig") if a.docs else pd.DataFrame()
    filas = []
    for i, r in enumerate(web.to_dict("records"), 1):
        term = ", ".join(k[2:] for k in r if k.startswith("n_") and r[k])
        filas.append(
            {
                "id": f"W{i:03d}",
                "gore": r["gore"],
                "fuente": "Sitio web",
                "ubicacion": r["url"],
                "fecha_o_pagina": r.get("fecha", ""),
                "terminos": term,
                "ods_nombrados": r.get("ods_nombrados", ""),
                "contexto": r.get("contexto", ""),
            }
        )
    man = pd.read_csv(a.manual, encoding="utf-8-sig") if getattr(a, "manual", None) else pd.DataFrame()
    vistas = {str(f["ubicacion"]).rstrip("/") for f in filas}
    k = 0
    for r in man.to_dict("records"):
        if str(r["url"]).rstrip("/") in vistas:
            continue  # el código ya la había encontrado
        k += 1
        filas.append(
            {
                "id": f"M{k:03d}",
                "gore": r["gore"],
                "fuente": "Sitio web",
                "ubicacion": r["url"],
                "fecha_o_pagina": r.get("fecha", ""),
                "terminos": r.get("terminos", ""),
                "ods_nombrados": "",
                "contexto": f"(búsqueda manual: {r.get('como_se_encontro', '')}) {r.get('titulo', '')}",
            }
        )
    for i, r in enumerate(docs.to_dict("records"), 1):
        filas.append(
            {
                "id": f"D{i:03d}",
                "gore": r["gore"],
                "fuente": r["fuente"],
                "ubicacion": r["id"],
                "fecha_o_pagina": f"p. {r['pagina']}",
                "terminos": r["terminos"],
                "ods_nombrados": r.get("ods_nombrados", ""),
                "contexto": r.get("contexto", ""),
            }
        )
    from openpyxl import Workbook

    wb = Workbook()
    m = wb.active
    m.title = "Menciones"
    cols = [
        "ID",
        "GORE",
        "Fuente",
        "Ubicación (URL o documento)",
        "Fecha / página",
        "Términos",
        "ODS nombrados",
        "Contexto",
        "Sugerencia automática",
        "Tipo (codificador 1)",
        "Justificación cod. 1",
        "Tipo (codificador 2)",
        "Justificación cod. 2",
        "Acuerdo",
        "Consenso final",
    ]
    encabezado(m, 1, cols, [7, 16, 13, 40, 12, 22, 12, 80, 18, 14, 30, 14, 30, 9, 14])
    dv = DataValidation(type="list", formula1='"Retórica,Sustantiva"', allow_blank=True)
    m.add_data_validation(dv)
    for r, f in enumerate(filas, 2):
        vals = [
            f["id"],
            f["gore"],
            f["fuente"],
            f["ubicacion"],
            f["fecha_o_pagina"],
            f["terminos"],
            f["ods_nombrados"],
            f["contexto"],
            sugerencia(f),
            None,
            None,
            None,
            None,
            f'=IF(OR(J{r}="",L{r}=""),"",IF(J{r}=L{r},1,0))',
            None,
        ]
        for c, v in enumerate(vals, 1):
            x = m.cell(r, c, v)
            x.font = F
            x.alignment = Alignment(wrap_text=c in (4, 8, 11, 13), vertical="top")
        for c in (10, 11, 12, 13, 15):
            m.cell(r, c).fill = AMARILLO
        dv.add(f"J{r}")
        dv.add(f"L{r}")
        dv.add(f"O{r}")
    ult = len(filas) + 1
    m.cell(ult + 2, 8, "Acuerdo observado:").font = FB
    m.cell(ult + 2, 14, f'=IFERROR(AVERAGE(N2:N{ult}),"")').number_format = "0%"
    m.freeze_panes = "B2"
    m.auto_filter.ref = f"A1:O{ult}"

    # Documentos: una fila por ERD y cuenta pública; la clasificación del documento completo
    dc = wb.create_sheet("Documentos")
    encabezado(
        dc,
        1,
        [
            "Código",
            "GORE",
            "Fuente",
            "Archivo / origen",
            "Estado automático",
            "Páginas con término",
            "ODS nombrados",
            "Clasificación (cod. 1)",
            "Clasificación (cod. 2)",
            "Acuerdo",
            "Consenso final",
            "Nota",
        ],
        [8, 16, 14, 30, 22, 10, 16, 16, 16, 9, 16, 40],
    )
    dv3 = DataValidation(type="list", formula1='"' + ",".join(CLASES_DOC) + '"', allow_blank=True)
    dc.add_data_validation(dv3)
    dd = pd.read_csv(a.documentos, encoding="utf-8-sig", dtype={"codigo": str}) if a.documentos else pd.DataFrame()
    for r, f in enumerate(dd.to_dict("records"), 2):
        nd = f.get("estado") in ("No disponible", "Solo video")
        vals = [
            f["codigo"],
            f["gore"],
            f["fuente"],
            f.get("archivo") if isinstance(f.get("archivo"), str) else f.get("origen"),
            f.get("estado"),
            f.get("paginas_con_termino_nucleo", ""),
            f.get("ods_nombrados", ""),
            "No disponible" if nd else None,
            "No disponible" if nd else None,
            f'=IF(OR(H{r}="",I{r}=""),"",IF(H{r}=I{r},1,0))',
            "No disponible" if nd else None,
            None,
        ]
        for c, v in enumerate(vals, 1):
            x = dc.cell(r, c, None if (isinstance(v, float) and pd.isna(v)) else v)
            x.font = F
        for c in (8, 9, 11, 12):
            dc.cell(r, c).fill = AMARILLO
        dv3.add(f"H{r}")
        dv3.add(f"I{r}")
        dv3.add(f"K{r}")
    dc.freeze_panes = "D2"

    # Mapa 17 ODS: tres filas por GORE (cod. 1, cod. 2, consenso), con el nivel de integración de la ERD
    mp = wb.create_sheet("Mapa_17_ODS")
    encabezado(
        mp,
        1,
        ["Código", "GORE", "Fila", "Documento ERD", "Nivel (0-3)"]
        + [f"ODS {n}" for n in range(1, 18)]
        + ["Candidatos detectados", "Evidencia / página"],
        [8, 16, 10, 26, 9] + [6] * 17 + [22, 40],
    )
    dv2 = DataValidation(type="list", formula1='"0,1,2"', allow_blank=True)
    mp.add_data_validation(dv2)
    dvn = DataValidation(type="list", formula1='"0,1,2,3"', allow_blank=True)
    mp.add_data_validation(dvn)
    erd = docs[docs.fuente == "ERD"] if not docs.empty else pd.DataFrame(columns=["gore", "ods_nombrados"])
    r = 2
    for g in cargar_gores():
        cand = sorted(
            {
                int(x)
                for s in erd[erd.gore == g["gore"]].ods_nombrados.fillna("").astype(str)
                for x in s.split()
                if x.isdigit()
            }
        )
        for fila in ("cod. 1", "cod. 2", "consenso"):
            mp.cell(r, 1, g["codigo"]).font = F
            mp.cell(r, 2, g["gore"]).font = F
            mp.cell(r, 3, fila).font = FB if fila == "consenso" else F
            for c in [4, 5] + list(range(6, 23)) + [24]:
                mp.cell(r, c).fill = AMARILLO
            dvn.add(mp.cell(r, 5))
            for n in range(1, 18):
                dv2.add(mp.cell(r, 5 + n))
            mp.cell(r, 23, " ".join(map(str, cand))).font = F
            r += 1
    mp.freeze_panes = "F2"

    # Muestra para repetir la búsqueda web (verificación de conteos)
    mu = wb.create_sheet("Muestra_verificacion")
    encabezado(
        mu,
        1,
        ["GORE", "Término", "Conteo cod. 1", "Conteo cod. 2", "Coincide (±1)", "Nombre cod. 2", "Fecha"],
        [18, 34, 12, 12, 12, 18, 12],
    )
    if a.busquedas:
        b = pd.read_csv(a.busquedas, encoding="utf-8-sig")
        random.seed(2026)
        muestra = random.sample(sorted(b.gore.unique()), min(5, b.gore.nunique()))
        r = 2
        for f in b[b.gore.isin(muestra)].to_dict("records"):
            mu.cell(r, 1, f["gore"])
            mu.cell(r, 2, f["termino"])
            mu.cell(r, 3, f["con_termino_verificado"])
            mu.cell(r, 5, f'=IF(D{r}="","",IF(ABS(D{r}-C{r})<=1,1,0))')
            for c in (4, 6, 7):
                mu.cell(r, c).fill = AMARILLO
            r += 1
        mu.cell(r + 1, 2, "% coincidentes (criterio ≥ 90 %):").font = FB
        mu.cell(r + 1, 5, f'=IFERROR(AVERAGE(E2:E{r - 1}),"")').number_format = "0%"
    out = SALIDAS / f"codificacion_ods_{a.etiqueta}.xlsx"
    wb.save(out)
    print(f"Plantilla: {out} ({len(filas)} menciones)")


def fiabilidad(a):
    """Comando «kappa»: acuerdo y κ de Cohen entre los dos codificadores de una planilla."""
    m = pd.read_excel(a.archivo, sheet_name="Menciones")
    m = m[m["ID"].notna()]
    par = m.dropna(subset=["Tipo (codificador 1)", "Tipo (codificador 2)"])
    po, k = kappa(par["Tipo (codificador 1)"], par["Tipo (codificador 2)"])
    print(f"Tipo de mención: n = {len(par)}; acuerdo = {po:.1%}; κ = {k:.2f}")
    print(pd.crosstab(par["Tipo (codificador 1)"], par["Tipo (codificador 2)"]))
    pend = m[m["Tipo (codificador 2)"].isna()]["ID"].tolist()
    if pend:
        print("Sin segunda codificación:", pend)
    mp = pd.read_excel(a.archivo, sheet_name="Mapa_17_ODS")
    ods = [f"ODS {n}" for n in range(1, 18)]
    c1 = mp[mp.Fila == "cod. 1"].set_index("GORE")[ods]
    c2 = mp[mp.Fila == "cod. 2"].set_index("GORE")[ods]
    llenos = c1.notna().all(axis=1) & c2.reindex(c1.index).notna().all(axis=1)
    if llenos.any():
        po, k = kappa(c1[llenos].values.ravel(), c2.reindex(c1.index)[llenos].values.ravel())
        print(
            f"Mapa 17 ODS: {int(llenos.sum())} ERD × 17 = {int(llenos.sum()) * 17} celdas; acuerdo = {po:.1%}; κ = {k:.2f}"
        )
    dc = pd.read_excel(a.archivo, sheet_name="Documentos").dropna(
        subset=["Clasificación (cod. 1)", "Clasificación (cod. 2)"]
    )
    if len(dc):
        po, k = kappa(dc["Clasificación (cod. 1)"], dc["Clasificación (cod. 2)"])
        print(f"Clasificación de documentos: n = {len(dc)}; acuerdo = {po:.1%}; κ = {k:.2f}")
    # resumen por GORE y fuente con el consenso (o, si falta, el codificador 1)
    m["final"] = m["Consenso final"].fillna(m["Tipo (codificador 1)"])
    niv = (
        m.assign(r=m.final.map({"Sustantiva": 2, "Retórica": 1}).fillna(0))
        .groupby(["GORE", "Fuente"])
        .r.max()
        .map({2: "Sustantiva", 1: "Retórica", 0: "Sin codificar"})
        .unstack()
    )
    niv.to_csv(SALIDAS / f"ods_resumen_{hoy()}.csv", encoding="utf-8-sig")
    print(niv.fillna("Sin mención").to_string())


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="accion", required=True)
    p = sub.add_parser("plantilla")
    p.add_argument("--web")
    p.add_argument("--docs", help="CSV de pasajes (ods_pasajes_FECHA.csv)")
    p.add_argument("--documentos", help="CSV de documentos (ods_documentos_FECHA.csv)")
    p.add_argument("--busquedas")
    p.add_argument(
        "--manual", help="CSV de menciones halladas a mano (ods_menciones_manuales_FECHA.csv; ver revision_manual.py)"
    )
    p.add_argument("--etiqueta", default=hoy())
    k = sub.add_parser("kappa")
    k.add_argument("archivo")
    a = ap.parse_args()
    plantilla(a) if a.accion == "plantilla" else fiabilidad(a)


if __name__ == "__main__":
    main()
