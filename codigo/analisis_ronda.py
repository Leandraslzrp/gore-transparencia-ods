"""
Análisis estadístico de una ronda de medición (cualquier ronda cerrada con cerrar_ronda.py).

  python codigo/analisis_ronda.py 2027-1                       # compara con la ronda anterior, si existe
  python codigo/analisis_ronda.py 2027-1 --comparar 2026-1
  python codigo/analisis_ronda.py 2027-1 --cplt-anio 2027      # por defecto, el año más reciente de datos/referencia_cplt.csv

Salida: resultados/<ronda>/analisis_<ronda>.xlsx (una hoja por análisis + Resumen + Fuentes) y resumen_<ronda>.csv.
Además actualiza resultados/serie_rondas.csv con los indicadores principales de TODAS las rondas cerradas, para
seguir la evolución en el tiempo.

Análisis incluidos:
  · CPLT: medias por año, Friedman + W de Kendall (≥ 3 años), Wilcoxon y Spearman entre años
  · Portal: índice por GORE, descriptivos, sensibilidad a la regla de puntaje, disponibilidad por categoría
  · Convergencia con el CPLT: Spearman con IC bootstrap, Wilcoxon, Bland-Altman, κ del umbral 90 %
  · Cambio respecto de la ronda anterior: celdas idénticas, κ, Spearman, Wilcoxon; cambios ODS por GORE
  · ODS: proporciones con IC de Wilson; nivel de la ERD vs transparencia (Spearman)
  · Fiabilidad entre codificadores y correlación mínima detectable
"""

import argparse
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

from comun import RAIZ, REF_CPLT
from ronda import CATS, carpeta, indice, kappa, leer, rondas

NOMBRES = {
    "A": "Estructura orgánica",
    "B": "Facultades, funciones y atribuciones",
    "C": "Marco normativo",
    "D1": "Personal de planta",
    "D2": "Personal a contrata",
    "D3": "Personal a honorarios",
    "E": "Compras y adquisiciones",
    "F": "Transferencias",
    "G": "Actos con efectos sobre terceros",
    "H": "Trámites",
    "I": "Subsidios y beneficios",
    "J": "Participación ciudadana",
    "K": "Ejecución presupuestaria",
    "L": "Auditorías",
    "M": "Vínculos institucionales",
}
REGLAS = {
    "base (OK=1, EX=1)": ({"OK": 1, "EX": 1}, False),
    "EX = 0": ({"OK": 1}, False),
    "EX = 0,5": ({"OK": 1, "EX": 0.5}, False),
    "sin ER": ({"OK": 1, "EX": 1}, True),
}
SEMILLA = 20260923
RESULTADOS = RAIZ / "resultados"


def wilson(k, n, z=1.959964):
    """Proporción k/n con su intervalo de confianza de Wilson (adecuado para muestras pequeñas)."""
    if n == 0:
        return float("nan"), float("nan"), float("nan")
    p = k / n
    den = 1 + z**2 / n
    centro = (p + z**2 / (2 * n)) / den
    margen = z * np.sqrt(p * (1 - p) / n + z**2 / (4 * n**2)) / den
    return p, max(0.0, centro - margen), min(1.0, centro + margen)


def spearman_boot(x, y, reps=5000):
    """Intervalo de confianza del 95 % de ρ de Spearman por remuestreo (bootstrap, semilla fija)."""
    rng = np.random.default_rng(SEMILLA)
    n, boot = len(x), []
    for _ in range(reps):
        i = rng.integers(0, n, n)
        if len(set(x[i])) > 1 and len(set(y[i])) > 1:
            boot.append(stats.spearmanr(x[i], y[i]).statistic)
    return np.percentile(boot, [2.5, 97.5])


def r_efecto(p, n):
    """Tamaño del efecto r de Wilcoxon a partir del valor p y del número de pares."""
    return stats.norm.isf(p / 2) / np.sqrt(n)


def wilcoxon_p(a, b):
    """Valor p de Wilcoxon; NaN si las dos series son idénticas (la prueba no se puede calcular)."""
    return stats.wilcoxon(a, b).pvalue if np.any(np.asarray(a) != np.asarray(b)) else float("nan")


def ronda_anterior(r):
    """Nombre de la ronda previa a «r», o None si es la primera."""
    todas = rondas()
    return todas[todas.index(r) - 1] if r in todas and todas.index(r) > 0 else None


def indicadores_ods(men, docs, mapa, n_gore):
    """Proporciones de divulgación de los ODS (web, ERD y cuentas públicas) con sus intervalos de Wilson."""
    web = men[men.fuente == "Sitio web"]
    cp = docs[docs.fuente == "Cuenta pública"].copy()
    cp["r"] = cp.consenso.map({"Sustantiva": 3, "Retórica": 2, "Sin mención": 1, "No disponible": 0}).fillna(0)
    cp = cp.sort_values("r", ascending=False).drop_duplicates("codigo")
    cp_disp = cp[cp.consenso != "No disponible"]
    niv = mapa.nivel
    return [
        ("GORE con mención ODS en su sitio web", web.codigo.nunique(), n_gore),
        ("Menciones web retóricas", int((web.consenso == "Retórica").sum()), len(web)),
        ("ERD que mencionan los ODS (nivel ≥ 1)", int((niv >= 1).sum()), int(niv.notna().sum())),
        ("ERD con vínculo sustantivo (nivel ≥ 2)", int((niv >= 2).sum()), int(niv.notna().sum())),
        ("ERD operacionalizadas (nivel 3)", int((niv == 3).sum()), int(niv.notna().sum())),
        ("Cuentas públicas con vínculo sustantivo", int((cp_disp.consenso == "Sustantiva").sum()), len(cp_disp)),
        (
            "Cuentas públicas con alguna mención",
            int(cp_disp.consenso.isin(["Retórica", "Sustantiva"]).sum()),
            len(cp_disp),
        ),
        ("ERD que vinculan el ODS 16", int((mapa.ods_16 == 2).sum()), len(mapa)),
    ]


def analizar(r, previa, anio):
    """Calcula todos los indicadores de una ronda y la compara con la anterior; devuelve hojas y resumen."""
    hojas, resumen = {}, []
    m = leer(r, "portal").set_index("codigo")
    men, docs, mapa = leer(r, "menciones"), leer(r, "documentos"), leer(r, "mapa")
    cplt = pd.read_csv(REF_CPLT, encoding="utf-8-sig", dtype={"codigo": str}).set_index("codigo")
    anios = [int(c[5:]) for c in cplt.columns if c.startswith("cplt_") and cplt[c].notna().all()]
    anio = anio or max(anios)
    cod = list(cplt.index)
    m = m.reindex(cod)
    if m[CATS].isna().any().any():
        raise SystemExit("La matriz del Portal no tiene los 16 GORE de referencia_cplt.csv.")
    n = len(cod)
    idx = {nombre: indice(m, *reglas).values for nombre, reglas in REGLAS.items()}
    base = idx["base (OK=1, EX=1)"]

    # Índices por GORE
    t = cplt.copy()
    t[f"Índice Portal ({r})"] = base.round(1)
    t["Nivel ERD"] = mapa.set_index("codigo").nivel.reindex(cod).values
    hojas["Indices_por_GORE"] = t.reset_index()

    # CPLT: evolución
    filas = []
    cols = [f"cplt_{a}" for a in sorted(anios)]
    if len(cols) >= 3:
        fr = stats.friedmanchisquare(*[cplt[c] for c in cols])
        W = fr.statistic / (n * (len(cols) - 1))
        filas.append(
            {
                "Prueba": f"Friedman ({', '.join(c[5:] for c in cols)})",
                "Estadístico": fr.statistic,
                "p": fr.pvalue,
                "Efecto": W,
                "Medida": "W de Kendall",
            }
        )
        resumen.append(("Friedman CPLT", f"χ²({len(cols) - 1}) = {fr.statistic:.1f}; p = {fr.pvalue:.3f}; W = {W:.2f}"))
    for a, b in zip(cols, cols[1:]):
        p = wilcoxon_p(cplt[b], cplt[a])
        dif = cplt[b] - cplt[a]
        filas.append(
            {
                "Prueba": f"Wilcoxon {a[5:]} → {b[5:]}",
                "Estadístico": "",
                "p": p,
                "Efecto": r_efecto(p, n),
                "Medida": "r = Z/√n",
                "Mejoran": int((dif > 0).sum()),
                "Empeoran": int((dif < 0).sum()),
                "Δ media (p.p.)": dif.mean(),
            }
        )
        s = stats.spearmanr(cplt[a], cplt[b])
        filas.append(
            {"Prueba": f"Spearman rankings {a[5:]} – {b[5:]}", "Estadístico": s.statistic, "p": s.pvalue, "Medida": "ρ"}
        )
    hojas["CPLT_evolucion"] = pd.DataFrame(filas)
    resumen.insert(
        0, (f"CPLT promedio {' / '.join(c[5:] for c in cols)} (%)", " / ".join(f"{cplt[c].mean():.1f}" for c in cols))
    )

    # Portal: descriptivos, sensibilidad, categorías
    hojas["Sensibilidad"] = pd.DataFrame(
        [
            {
                "Regla": k,
                "Media": v.mean(),
                "Mediana": np.median(v),
                "DE": v.std(ddof=1),
                "Mín.": v.min(),
                "Máx.": v.max(),
                "GORE ≥ 90 %": int((v >= 90).sum()),
            }
            for k, v in idx.items()
        ]
    )
    resumen.append(
        (
            f"Índice medio del Portal ({r})",
            f"{base.mean():.1f} % (DE {base.std(ddof=1):.1f}); EX = 0: {idx['EX = 0'].mean():.1f} %",
        )
    )
    filas = []
    for c in CATS:
        vc = m[c].value_counts()
        filas.append(
            {
                "Cat.": c,
                "Categoría": NOMBRES[c],
                **{e: int(vc.get(e, 0)) for e in ("OK", "EX", "SD", "ER")},
                "% disponible": (vc.get("OK", 0) + vc.get("EX", 0)) / n * 100,
            }
        )
    hojas["Categorias"] = pd.DataFrame(filas)

    # Convergencia con el CPLT
    x = cplt[f"cplt_{anio}"].values
    s = stats.spearmanr(x, base)
    lo, hi = spearman_boot(x, base)
    dif = base - x
    po, k90 = kappa(x >= 90, base >= 90)
    conv = {
        "Comparación": f"Índice Portal ({r}) vs CPLT {anio}",
        "Spearman ρ": s.statistic,
        "Spearman p": s.pvalue,
        "IC 95 % bootstrap inf.": lo,
        "IC 95 % bootstrap sup.": hi,
        "Wilcoxon p": wilcoxon_p(base, x),
        "Bland-Altman sesgo (p.p.)": dif.mean(),
        "Límite inferior": dif.mean() - 1.96 * dif.std(ddof=1),
        "Límite superior": dif.mean() + 1.96 * dif.std(ddof=1),
        "Acuerdo ≥ 90 %": po,
        "κ ≥ 90 %": k90,
    }
    hojas["Convergencia_CPLT"] = pd.DataFrame([conv])
    resumen += [
        (
            f"Índice Portal vs CPLT {anio}",
            f"ρ = {conv['Spearman ρ']:.2f} (IC {lo:.2f} a {hi:.2f}); "
            f"Wilcoxon p = {conv['Wilcoxon p']:.3f}; κ ≥ 90 % = {k90:.2f}",
        ),
        (
            "Bland-Altman",
            f"sesgo {conv['Bland-Altman sesgo (p.p.)']:.1f} p.p.; límites {conv['Límite inferior']:.1f} a {conv['Límite superior']:.1f}",
        ),
    ]

    # Cambio respecto de la ronda anterior
    if previa:
        mp = leer(previa, "portal").set_index("codigo").reindex(cod)
        ea, eb = mp[CATS].values.ravel(), m[CATS].values.ravel()
        ba, bb = np.isin(ea, ["OK", "EX"]).astype(int), np.isin(eb, ["OK", "EX"]).astype(int)
        po, k = kappa(ba, bb)
        ia = indice(mp).values
        cambio = {
            "Rondas": f"{previa} → {r}",
            "Estado idéntico": f"{int((ea == eb).sum())} de {len(ea)}",
            "0/1 idéntico": f"{int((ba == bb).sum())} de {len(ea)}",
            "κ (0/1)": k,
            "Spearman ρ índices": stats.spearmanr(ia, base).statistic,
            "Wilcoxon p": wilcoxon_p(base, ia),
            "Δ índice medio (p.p.)": base.mean() - ia.mean(),
            "GORE que mejoran": int((base > ia).sum()),
            "GORE que empeoran": int((base < ia).sum()),
        }
        hojas["Cambio_ronda_anterior"] = pd.DataFrame([cambio])
        resumen.append(
            (
                f"Cambio {previa} → {r} (Portal)",
                f"Δ = {cambio['Δ índice medio (p.p.)']:+.1f} p.p.; κ = {k:.2f}; "
                f"Wilcoxon p = {cambio['Wilcoxon p']:.3f}",
            )
        )
        ra, rb = leer(previa, "resumen").set_index("codigo"), leer(r, "resumen").set_index("codigo")
        filas = []
        for c in cod:
            filas.append(
                {
                    "codigo": c,
                    "gore": rb.loc[c, "gore"],
                    "Nivel ERD antes": ra.loc[c, "erd_nivel"],
                    "Nivel ERD ahora": rb.loc[c, "erd_nivel"],
                    "Cuenta pública antes": ra.loc[c, "cuenta_publica"],
                    "Cuenta pública ahora": rb.loc[c, "cuenta_publica"],
                    "Menciones web antes": ra.loc[c, "web_retoricas"] + ra.loc[c, "web_sustantivas"],
                    "Menciones web ahora": rb.loc[c, "web_retoricas"] + rb.loc[c, "web_sustantivas"],
                }
            )
        hojas["Cambio_ODS_por_GORE"] = pd.DataFrame(filas)

    # ODS
    filas = []
    for nombre, k, nn in indicadores_ods(men, docs, mapa, n):
        p, lo, hi = wilson(k, nn)
        filas.append(
            {
                "Indicador": nombre,
                "k": k,
                "n": nn,
                "Proporción": p,
                "IC 95 % inf. (Wilson)": lo,
                "IC 95 % sup. (Wilson)": hi,
            }
        )
        resumen.append((nombre, f"{k} de {nn} ({p * 100:.0f} %; IC 95 %: {lo * 100:.0f}–{hi * 100:.0f} %)"))
    hojas["ODS_proporciones"] = pd.DataFrame(filas)
    hojas["ODS_mapa"] = mapa
    niv = mapa.set_index("codigo").nivel.reindex(cod).values.astype(float)
    filas = []
    for nombre, v in [(f"CPLT {anio}", x), (f"Índice Portal ({r})", base)]:
        s = stats.spearmanr(niv, v)
        filas.append({"Nivel ERD vs": nombre, "ρ": s.statistic, "p": s.pvalue, "n": n})
    hojas["ODS_vs_transparencia"] = pd.DataFrame(filas)
    resumen.append((f"Nivel ERD vs CPLT {anio}", f"ρ = {filas[0]['ρ']:.2f}; p = {filas[0]['p']:.2f}"))

    # Fiabilidad y potencia
    fia = leer(r, "fiabilidad")
    hojas["Fiabilidad"] = fia
    for f in fia.to_dict("records"):
        if pd.notna(f["kappa"]):
            resumen.append(
                (f"κ {f['medida']}", f"{f['kappa']:.2f} (n = {int(f['n'])}; acuerdo {f['acuerdo'] * 100:.0f} %)")
            )
        elif pd.notna(f["acuerdo"]):
            resumen.append((f["medida"], f"{round(f['acuerdo'] * f['n'])} de {int(f['n'])}"))
    r_min = np.tanh((stats.norm.isf(0.025) + stats.norm.isf(0.20)) / np.sqrt(n - 3))
    hojas["Potencia"] = pd.DataFrame(
        [{"Medida": f"Correlación mínima detectable (α = 0,05; potencia 80 %; n = {n})", "Valor": round(r_min, 2)}]
    )
    resumen.append((f"Correlación mínima detectable (n = {n})", f"|ρ| ≈ {r_min:.2f}"))
    return hojas, pd.DataFrame(resumen, columns=["Indicador", "Valor"]), anio


def fila_serie(r):
    """Fila de resultados/serie_rondas.csv con los indicadores principales de una ronda."""
    m = leer(r, "portal")
    man = leer(r, "manifiesto")
    ods = {
        k: (a, b) for k, a, b in indicadores_ods(leer(r, "menciones"), leer(r, "documentos"), leer(r, "mapa"), len(m))
    }
    f = {
        "ronda": r,
        "fecha_scraping": man.get("fecha_scraping", ""),
        "fecha_ods": man.get("fecha_ods", ""),
        "indice_portal_medio": round(indice(m).mean(), 1),
    }
    for k, (a, b) in ods.items():
        f[k] = f"{a}/{b}"
    return f


def escribir(hojas, resumen, ruta, fuentes, titulo):
    """Guarda el libro Excel de resultados con una hoja de resumen, una por análisis y una de fuentes."""
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        resumen.to_excel(xw, sheet_name="Resumen", index=False, startrow=3)
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre[:31], index=False)
        pd.DataFrame(fuentes, columns=["Dato", "Archivo"]).to_excel(xw, sheet_name="Fuentes", index=False)
    wb = load_workbook(ruta)
    ws = wb["Resumen"]
    ws["A1"] = titulo
    ws["A2"] = f"Generado el {date.today():%d-%m-%Y} por analisis_ronda.py. La hoja Fuentes indica los archivos usados."
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    azul = PatternFill("solid", fgColor="DDEBF7")
    for w in wb.worksheets:
        fe = 4 if w.title == "Resumen" else 1
        for c in w[fe]:
            c.font = Font(name="Arial", size=10, bold=True)
            c.fill = azul
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for fila in w.iter_rows(min_row=fe + 1):
            for c in fila:
                if not (c.font and c.font.b):
                    c.font = Font(name="Arial", size=10)
                if isinstance(c.value, float):
                    c.number_format = "0.000" if abs(c.value) < 1 else "0.0"
        for i, col in enumerate(w.columns, 1):
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            w.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 70)
        w.freeze_panes = w.cell(fe + 1, 1)
    wb.save(ruta)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ronda")
    ap.add_argument("--comparar", help="ronda con la que comparar (por defecto, la anterior en orden alfabético)")
    ap.add_argument("--sin-comparar", action="store_true")
    ap.add_argument("--cplt-anio", type=int)
    a = ap.parse_args()
    carpeta(a.ronda)
    previa = None if a.sin_comparar else (a.comparar or ronda_anterior(a.ronda))
    hojas, resumen, anio = analizar(a.ronda, previa, a.cplt_anio)
    out = RESULTADOS / a.ronda
    out.mkdir(parents=True, exist_ok=True)
    ruta = out / f"analisis_{a.ronda}.xlsx"
    fuentes = [
        ("Ronda", f"mediciones/{a.ronda}/ (ver manifiesto.json)"),
        ("CPLT", f"datos/referencia_cplt.csv (año {anio})"),
        ("Ronda comparada", f"mediciones/{previa}/" if previa else "—"),
    ]
    escribir(hojas, resumen, ruta, fuentes, f"Análisis de la ronda {a.ronda}")
    resumen.to_csv(out / f"resumen_{a.ronda}.csv", index=False, encoding="utf-8")
    serie = pd.DataFrame([fila_serie(x) for x in rondas()])
    serie.to_csv(RESULTADOS / "serie_rondas.csv", index=False, encoding="utf-8")
    print(resumen.to_string(index=False))
    print(f"\nListo: {ruta.relative_to(RAIZ)} · serie actualizada: resultados/serie_rondas.csv ({len(serie)} rondas)")


if __name__ == "__main__":
    main()
