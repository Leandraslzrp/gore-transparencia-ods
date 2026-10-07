"""
Pruebas automáticas sin internet. Ejecutar desde la carpeta del proyecto:
  python -m unittest discover -s tests -v        (o: python ejecutar_todo.py pruebas)
Comprueban que (1) las reglas de búsqueda y de cálculo no cambiaron sin querer y (2) el flujo completo de una ronda
(revisión manual → codificación → cierre → análisis) funciona con los datos de ejemplo de tests/datos/
(salidas reales de la recolección automática del 29-09-2026).
"""

import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from urllib.robotparser import RobotFileParser

RAIZ = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RAIZ / "codigo"))

import pandas as pd  # noqa: E402

import comun  # noqa: E402
from analisis_ronda import wilson  # noqa: E402
from ods_terminos import contar, ods_mencionados, tiene_nucleo  # noqa: E402
from ronda import _orden, codigo_de, indice, kappa  # noqa: E402


class Terminos(unittest.TestCase):
    def test_conteo(self):
        t = (
            "La ERD se alinea con los Objetivos de Desarrollo Sostenible (ODS) de la Agenda 2030. "
            "El eje 4 contribuye al ODS 16. Métodos y todos no cuentan; ods en minúscula tampoco."
        )
        c = contar(t)
        self.assertEqual(c["ODS"], 2)
        self.assertEqual(c["ODS 16"], 1)
        self.assertEqual(c["Agenda 2030"], 1)
        self.assertEqual(c["Objetivos de Desarrollo Sostenible"], 1)

    def test_falsos_positivos(self):
        self.assertFalse(tiene_nucleo("Métodos, todos, PERIODOS y la ruta /ODS/ no son menciones."))

    def test_ods_mencionados(self):
        o = ods_mencionados(
            "Contribuye al ODS 11 y al Objetivo de Desarrollo Sostenible 6. Tabla: 16. Paz, justicia; meta 16.6. "
            "Igualdad de género."
        )
        self.assertEqual(o["todos"], [5, 6, 11, 16])
        self.assertEqual(o["metas"], ["16.6"])


class Robots(unittest.TestCase):
    def setUp(self):
        self.prev = dict(comun._ROBOTS)
        comun.configurar_robots(False)

    def tearDown(self):
        comun._ROBOTS.clear()
        comun._ROBOTS.update(self.prev)
        comun.configurar_robots(False)

    def _sitio(self, reglas):
        rp = RobotFileParser()
        rp.parse(reglas.splitlines())
        comun._ROBOTS["https://ejemplo.cl"] = rp

    def test_prohibido_para_todos(self):
        self._sitio("User-agent: *\nDisallow: /privado/")
        self.assertFalse(comun.robots_permite("https://ejemplo.cl/privado/x.pdf"))
        self.assertTrue(comun.robots_permite("https://ejemplo.cl/noticias/ods"))

    def test_prohibido_para_nuestro_agente(self):
        self._sitio("User-agent: InvestigacionTransparenciaGORE\nDisallow: /\n\nUser-agent: *\nAllow: /")
        self.assertFalse(comun.robots_permite("https://ejemplo.cl/cualquier"))

    def test_sin_robots_permite(self):
        comun._ROBOTS["https://ejemplo.cl"] = None  # 404 o error de lectura → permitido (convención de Google)
        self.assertTrue(comun.robots_permite("https://ejemplo.cl/a"))

    def test_ignorar(self):
        self._sitio("User-agent: *\nDisallow: /")
        comun.configurar_robots(True)
        self.assertTrue(comun.robots_permite("https://ejemplo.cl/a"))

    def test_sesion_bloquea_sin_conectarse(self):
        self._sitio("User-agent: *\nDisallow: /")
        with self.assertRaises(comun.BloqueadoPorRobots):
            comun.sesion().get("https://ejemplo.cl/a")


class Calculos(unittest.TestCase):
    def test_indice(self):
        m = pd.DataFrame(
            [["OK"] * 13 + ["EX", "SD"], ["ER"] * 15],
            columns=["A", "B", "C", "D1", "D2", "D3", "E", "F", "G", "H", "I", "J", "K", "L", "M"],
        )
        self.assertAlmostEqual(indice(m).iloc[0], 14 / 15 * 100)
        self.assertEqual(indice(m).iloc[1], 0)

    def test_wilson(self):
        p, lo, hi = wilson(13, 16)
        self.assertEqual((round(lo, 2), round(hi, 2)), (0.57, 0.93))

    def test_kappa(self):
        po, k = kappa(["a", "a", "b", "b"], ["a", "a", "b", "b"])
        self.assertEqual((po, k), (1.0, 1.0))
        po, k = kappa(["a", "b", "a", "b"], ["a", "a", "b", "b"])
        self.assertAlmostEqual(k, 0.0)

    def test_hash(self):
        self.assertEqual(comun.sha256_texto("ODS"), "85c9216cc1f071847636c42003dd7541adc7c73d168fff884e70d13e1bbe65db")
        self.assertEqual(len(comun.sha256_texto("x")), 64)

    def test_codigos_y_orden(self):
        self.assertEqual(codigo_de("Magallanes y de la Antártica Chilena"), "AB089")
        self.assertEqual(
            sorted(["2027-1", "2026-10", "2026-2", "piloto"], key=_orden), ["2026-2", "2026-10", "2027-1", "piloto"]
        )


class FlujoCompleto(unittest.TestCase):
    """Simula dos rondas de punta a punta en una copia temporal: revisión → codificación → cierre → análisis."""

    def test_flujo(self):
        e = "2026-09-29"
        with tempfile.TemporaryDirectory() as tmp:
            d = Path(tmp) / "repo"
            shutil.copytree(
                RAIZ,
                d,
                ignore=shutil.ignore_patterns(
                    "cache", "salidas", "resultados", ".git", ".venv", "__pycache__", "*.pdf"
                ),
            )
            (d / "salidas").mkdir()
            rp = RAIZ / "tests" / "datos"  # salidas de ejemplo de la recolección automática
            for f in rp.glob(f"*_{e}.csv"):
                shutil.copy(f, d / "salidas" / f.name)
            # los datos de ejemplo son anteriores a la columna sha256: se agrega una huella de prueba
            fd = d / "salidas" / f"ods_documentos_{e}.csv"
            dd = pd.read_csv(fd, encoding="utf-8-sig")
            dd["sha256"] = [comun.sha256_texto(f"prueba-{i}") for i in range(len(dd))]
            dd.to_csv(fd, index=False, encoding="utf-8-sig")
            py = lambda *a: subprocess.run([sys.executable, *a], cwd=d, check=True, capture_output=True, text=True)
            py(
                "codigo/ods_codificacion.py",
                "plantilla",
                "--web",
                f"salidas/ods_menciones_web_{e}.csv",
                "--docs",
                f"salidas/ods_pasajes_{e}.csv",
                "--documentos",
                f"salidas/ods_documentos_{e}.csv",
                "--busquedas",
                f"salidas/ods_busquedas_{e}.csv",
                "--etiqueta",
                e,
            )
            from openpyxl import load_workbook

            f = d / "salidas" / f"codificacion_ods_{e}.xlsx"
            wb = load_workbook(f)
            m = wb["Menciones"]
            for r in range(2, m.max_row + 1):
                if m.cell(r, 2).value:
                    m.cell(r, 10, "Retórica")
                    m.cell(r, 12, "Retórica")
            dc = wb["Documentos"]
            for r in range(2, dc.max_row + 1):
                if dc.cell(r, 1).value and not dc.cell(r, 8).value:
                    dc.cell(r, 8, "Sin mención")
                    dc.cell(r, 9, "Sin mención")
            mp = wb["Mapa_17_ODS"]
            for r in range(2, mp.max_row + 1):
                if mp.cell(r, 1).value:
                    for c in range(5, 23):
                        mp.cell(r, c, 0)
            wb.save(f)
            # revisión manual: se crea, se completa y se exporta
            py("codigo/revision_manual.py", "crear", "--etiqueta", e)
            rv = d / "salidas" / f"revision_manual_{e}.xlsx"
            wr = load_workbook(rv)
            wp = wr["Pendientes"]
            for r in range(2, wp.max_row + 1):
                wp.cell(r, 6, "Confirmado (sin cambios)")
                wp.cell(r, 8, "Prueba")
            wm = wr["Menciones_manuales"]
            for j, v in enumerate(
                [
                    "AB079",
                    "Coquimbo",
                    "https://ejemplo.cl/ods",
                    "Noticia",
                    "2025-01-01",
                    "ODS",
                    "Google",
                    "Prueba",
                    "2026-10-07",
                ],
                1,
            ):
                wm.cell(2, j, v)
            wr.save(rv)
            py("codigo/revision_manual.py", "exportar", "--etiqueta", e)
            self.assertTrue((d / "salidas" / f"ods_menciones_manuales_{e}.csv").exists())
            for r in ("2027-1", "2027-2"):  # dos rondas con los mismos datos, para probar la comparación
                py("codigo/cerrar_ronda.py", r, "--codificacion", str(f), "--etiqueta", e)
                out = py("codigo/analisis_ronda.py", r).stdout
            self.assertTrue((d / "mediciones" / "2027-2" / "crudos" / f"revision_manual_{e}.csv").exists())
            self.assertIn("Cambio 2027-1 → 2027-2", out)
            docs = pd.read_csv(d / "mediciones" / "2027-2" / "ods_documentos.csv")
            self.assertTrue(
                docs.sha256.notna().any()
            )  # la huella de cada PDF viene de la recolección, no de la planilla
            man = json.loads((d / "mediciones" / "2027-2" / "manifiesto.json").read_text(encoding="utf-8"))
            self.assertIn("portal_matriz.csv", man["archivos"])
            serie = pd.read_csv(d / "resultados" / "serie_rondas.csv")
            self.assertEqual(list(serie.ronda), ["2027-1", "2027-2"])
            self.assertEqual(serie.indice_portal_medio.iloc[-1], 89.2)  # recorrido del 29-09-2026
            ver = subprocess.run(
                [sys.executable, "codigo/cerrar_ronda.py", "--verificar", "2027-2"],
                cwd=d,
                capture_output=True,
                text=True,
            )
            self.assertIn("íntegra", ver.stdout)
            with self.assertRaises(subprocess.CalledProcessError):  # una ronda cerrada no se sobrescribe sin --forzar
                py("codigo/cerrar_ronda.py", "2027-2", "--codificacion", str(f), "--etiqueta", e)


class RevisionManual(unittest.TestCase):
    def test_no_disponible_exige_dos_revisores(self):
        import revision_manual

        self.assertIn("No disponible (confirmado por 2)", revision_manual.RESULTADOS)


if __name__ == "__main__":
    unittest.main()
