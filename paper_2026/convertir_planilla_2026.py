"""
Convierte los datos del paper 2026 (planilla consolidada y scraping de julio) al formato estándar de ronda.
Se ejecutó una vez para crear mediciones/2026-1/ (ronda del paper). Se conserva para documentar el origen de esos archivos.

  python paper_2026/convertir_planilla_2026.py
"""

import re
import shutil
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "codigo"))
from cerrar_ronda import manifiesto  # noqa: E402
from comun import CONFIG, RAIZ  # noqa: E402
from ronda import ARCHIVOS, carpeta, codigo_de, fiabilidad, indice, resumen_gore  # noqa: E402

R = "2026-1"  # ronda 1 de 2026 = datos del paper (Portal: julio; ODS: septiembre)
wb = load_workbook(RAIZ / "paper_2026" / "planilla_codificacion_ods.xlsx", data_only=True)
out = carpeta(R, crear=True)

m = pd.read_csv(RAIZ / "datos" / "referencia_scraping_2026-07-22.csv", encoding="utf-8-sig")
m.insert(0, "codigo", m.gore.map(codigo_de))
m["indice"] = indice(m).round(1)
m.to_csv(out / ARCHIVOS["portal"], index=False, encoding="utf-8")

men = []
for r in wb["Menciones"].iter_rows(min_row=5, values_only=True):
    if not r[0]:
        continue
    men.append(
        {
            "id": r[0],
            "codigo": codigo_de(r[1]),
            "gore": r[1],
            "fuente": r[2],
            "fecha_o_pagina": r[3],
            "titulo_o_documento": r[4],
            "url": r[5],
            "terminos": r[6],
            "sha256_texto": "",
            "cod1": r[7],
            "cod2": r[9],
            "consenso": r[11] or r[7],
            "justificacion": r[8],
        }
    )
pd.DataFrame(men).to_csv(out / ARCHIVOS["menciones"], index=False, encoding="utf-8")

docs = []
for r in wb["Documentos"].iter_rows(min_row=5, values_only=True):
    if not r[0] or r[1] not in ("Cuenta pública", "ERD"):
        continue
    docs.append(
        {
            "codigo": codigo_de(r[0]),
            "gore": r[0],
            "fuente": r[1],
            "documento": r[2],
            "url": r[3],
            "sha256": "",
            "paginas": r[4],
            "caracteres": r[5],
            "n_ODS": r[7],
            "n_Objetivos de Desarrollo Sostenible": r[8],
            "n_Agenda 2030": r[9],
            "n_ODS 16": r[10],
            "cod1": r[14],
            "cod2": "",
            "consenso": r[14],
            "nota": r[15],
        }
    )
pd.DataFrame(docs).to_csv(out / ARCHIVOS["documentos"], index=False, encoding="utf-8")

mapa = []
for r in wb["Mapa_17_ODS"].iter_rows(min_row=6, max_row=21, values_only=True):
    niv = re.match(r"(\d)", str(r[2] or ""))
    mapa.append(
        {
            "codigo": codigo_de(r[0]),
            "gore": r[0],
            "documento": r[1],
            "nivel": int(niv.group(1)) if niv else 0,
            **{f"ods_{k + 1}": int(r[3 + k] or 0) for k in range(17)},
        }
    )
mapa = pd.DataFrame(mapa)
mapa.to_csv(out / ARCHIVOS["mapa"], index=False, encoding="utf-8")

ver = []
for r in wb["Verificacion_Mapa"].iter_rows(min_row=5, values_only=True):
    if r[0] and r[2]:
        k = int(str(r[2]).split(".")[0])
        ver.append({"codigo": codigo_de(r[0]), "gore": r[0], "ods": k, "cod1": r[3], "cod2": r[4]})
pd.DataFrame(ver).to_csv(out / ARCHIVOS["mapa_cod"], index=False, encoding="utf-8")

ver2 = []
for r in wb["Verificacion"].iter_rows(min_row=5, values_only=True):
    if r[0] and r[1] == "Sitio web" and isinstance(r[3], (int, float)) and isinstance(r[4], (int, float)):
        ver2.append({"gore": r[0], "termino": r[2], "cod1": r[3], "cod2": r[4]})
men_df, docs_df = pd.DataFrame(men), pd.DataFrame(docs)
resumen_gore(men_df, docs_df, mapa).to_csv(out / ARCHIVOS["resumen"], index=False, encoding="utf-8")
fiabilidad(men_df, docs_df, pd.DataFrame(ver), pd.DataFrame(ver2)).to_csv(
    out / ARCHIVOS["fiabilidad"], index=False, encoding="utf-8"
)
shutil.copy(CONFIG / "gores.csv", out / "config_gores.csv")
manifiesto(
    R,
    out,
    {
        "origen": "Conversión de la planilla consolidada del paper (codificación 23–29 sep. 2026) y del "
        "scraping del 22 y 27 de julio de 2026 (datos/referencia_scraping_2026-07-22.csv).",
        "fecha_scraping": "22 y 27 de julio de 2026",
        "fecha_ods": "23 al 29 de septiembre de 2026",
    },
)
print(f"Ronda {R} creada en {out}")
