"""
Análisis estadístico del paper 2026 (CONGELADO: reproduce exactamente las cifras publicadas).
Para rondas nuevas use analisis_ronda.py.

Reúne en un solo libro todas las cifras del paper, calculadas desde los datos:
  · Índices del CPLT 2024–2026 ............ datos/referencia_cplt.csv
  · Scraping 1 (22-07-2026) ............... datos/referencia_scraping_2026-07-22.csv
  · Scraping 2 (25-09-2026) ............... datos/referencia_scraping_2026-09-25.csv
  · Scraping 3 (código Python, 29-09) ..... paper_2026/resultados/portal_matriz_2026-09-29.csv
  · Búsqueda ODS codificada por el equipo . paper_2026/planilla_codificacion_ods.xlsx
  · Búsqueda ODS hecha por el código ...... paper_2026/resultados/ods_menciones_web_*.csv y ods_documentos_*.csv

Uso (con el entorno activado, desde la raíz del repositorio):
  python ejecutar_todo.py paper
  python paper_2026/analisis_paper_2026.py --planilla "ruta/a/la/planilla.xlsx"
  python paper_2026/analisis_paper_2026.py --entrada salidas      # usar los recorridos más recientes del código
Salida: salidas/analisis_estadistico_FECHA.xlsx
Requiere pandas, numpy, scipy y openpyxl.
"""

import argparse
import glob
import os
import re
from datetime import date

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from scipy import stats

PAPER = os.path.dirname(os.path.abspath(__file__))
RAIZ = os.path.dirname(PAPER)
DATOS = os.path.join(RAIZ, "datos")
SALIDAS = os.path.join(RAIZ, "salidas")
PLANILLA = os.path.join(PAPER, "planilla_codificacion_ods.xlsx")
ENTRADA = os.path.join(PAPER, "resultados")  # salidas del código usadas en el paper (29-09-2026); --entrada la cambia
CATS = ["A", "B", "C", "D1", "D2", "D3", "E", "F", "G", "H", "I", "J", "K", "L", "M"]
NOMBRES = {
    "A": "Estructura orgánica",
    "B": "Facultades, funciones y atribuciones",
    "C": "Marco normativo",
    "D1": "Personal de planta",
    "D2": "Personal a contrata",
    "D3": "Personal a honorarios",
    "E": "Compras y adquisiciones",
    "F": "Transferencias",
    "G": "Actos con efectos sobre terceros",
    "H": "Trámites",
    "I": "Subsidios y beneficios",
    "J": "Participación ciudadana",
    "K": "Ejecución presupuestaria",
    "L": "Auditorías",
    "M": "Vínculos institucionales",
}
SEMILLA = 20260923


# ---------------------------------------------------------------- utilidades
def ultimo(patron):
    archivos = sorted(glob.glob(os.path.join(ENTRADA, patron)))
    return archivos[-1] if archivos else None


def kappa(a, b):
    a, b = pd.Series(list(a)), pd.Series(list(b))
    po = float((a == b).mean())
    pe = sum(float((a == c).mean()) * float((b == c).mean()) for c in set(a) | set(b))
    return po, (po - pe) / (1 - pe) if pe < 1 else float("nan")


def wilson(k, n, z=1.959964):
    if n == 0:
        return float("nan"), float("nan"), float("nan")
    p = k / n
    den = 1 + z**2 / n
    centro = (p + z**2 / (2 * n)) / den
    margen = z * np.sqrt(p * (1 - p) / n + z**2 / (4 * n**2)) / den
    return p, max(0.0, centro - margen), min(1.0, centro + margen)


def r_efecto(p, n):
    return stats.norm.isf(p / 2) / np.sqrt(n)


def indice(m, regla="base"):
    out = {}
    for g, fila in m.iterrows():
        v = fila[CATS].values
        if regla == "base":
            out[g] = np.isin(v, ["OK", "EX"]).mean() * 100
        elif regla == "EX = 0":
            out[g] = (v == "OK").mean() * 100
        elif regla == "EX = 0,5":
            out[g] = ((v == "OK").sum() + 0.5 * (v == "EX").sum()) / len(v) * 100
        elif regla == "sin ER":
            v2 = v[v != "ER"]
            out[g] = np.isin(v2, ["OK", "EX"]).mean() * 100 if len(v2) else np.nan
    return pd.Series(out)


def norm_url(u):
    u = str(u).strip().lower()
    u = re.sub(r"^https?://(www\.)?", "", u)
    return u.rstrip("/")


def spearman_boot(x, y, reps=5000):
    rng = np.random.default_rng(SEMILLA)
    n, boot = len(x), []
    for _ in range(reps):
        i = rng.integers(0, n, n)
        if len(set(x[i])) > 1 and len(set(y[i])) > 1:
            boot.append(stats.spearmanr(x[i], y[i]).statistic)
    return np.percentile(boot, [2.5, 97.5])


# ---------------------------------------------------------------- datos
def cargar(planilla):
    cplt = pd.read_csv(os.path.join(DATOS, "referencia_cplt.csv"))
    s1 = pd.read_csv(os.path.join(DATOS, "referencia_scraping_2026-07-22.csv")).set_index("gore")
    s2 = pd.read_csv(os.path.join(DATOS, "referencia_scraping_2026-09-25.csv")).set_index("gore")
    f3 = ultimo("portal_matriz_*.csv")
    if f3 is None:
        raise SystemExit(f"Falta portal_matriz_FECHA.csv en {ENTRADA}.")
    s3 = pd.read_csv(f3, encoding="utf-8-sig").set_index("gore")
    etiqueta3 = re.search(r"portal_matriz_(.+)\.csv", f3).group(1)
    orden = list(cplt.gore)
    s1, s2, s3 = (m.reindex(orden) for m in (s1, s2, s3))
    wb = load_workbook(planilla, data_only=True)
    return cplt, s1, s2, s3, etiqueta3, wb


def leer_planilla(wb):
    men = pd.DataFrame(
        [r for r in wb["Menciones"].iter_rows(min_row=5, values_only=True) if r[0]],
        columns=[
            "id",
            "gore",
            "fuente",
            "fecha",
            "titulo",
            "url",
            "terminos",
            "cod1",
            "justif",
            "cod2",
            "acuerdo",
            "consenso",
        ][:12],
    )
    docs = []
    for r in wb["Documentos"].iter_rows(min_row=5, values_only=True):
        if r[0] and r[1] in ("Cuenta pública", "ERD"):
            docs.append({"gore": r[0], "fuente": r[1], "clasificacion": r[14]})
    docs = pd.DataFrame(docs)
    rango = {"Sustantiva": 3, "Retórica": 2, "Sin mención": 1, "No disponible": 0}
    docs["rango"] = docs.clasificacion.map(rango).fillna(0)
    docs = docs.sort_values("rango", ascending=False).drop_duplicates(["gore", "fuente"])
    mapa = []
    for r in wb["Mapa_17_ODS"].iter_rows(min_row=6, max_row=21, values_only=True):
        nivel = str(r[2] or "")
        m = re.match(r"(\d)", nivel)
        mapa.append(
            {"gore": r[0], "peldano": int(m.group(1)) if m else np.nan, **{f"ODS {k + 1}": r[3 + k] for k in range(17)}}
        )
    mapa = pd.DataFrame(mapa)
    ver = []
    for r in wb["Verificacion"].iter_rows(min_row=4, max_row=80, values_only=True):
        if r[0] and r[1] in ("Sitio web", "Cuenta pública", "ERD") and isinstance(r[3], (int, float)):
            ver.append({"gore": r[0], "fuente": r[1], "termino": r[2], "cod1": r[3], "cod2": r[4]})
    return men, docs, mapa, pd.DataFrame(ver)


# ---------------------------------------------------------------- análisis
def analizar(cplt, s1, s2, s3, et3, men, docs, mapa, ver, web_cod, doc_cod):
    n = len(cplt)
    hojas, resumen = {}, []
    g = cplt.gore.values
    idx = {
        "Scraping 1 (22-07)": indice(s1).reindex(g).values,
        "Scraping 2 (25-09)": indice(s2).reindex(g).values,
        f"Scraping 3 ({et3}, código)": indice(s3).reindex(g).values,
    }

    # Índices por GORE
    t = cplt.copy()
    for k, v in idx.items():
        t[k] = v
    hojas["Indices_por_GORE"] = t

    # Descriptivos
    series = {
        "CPLT 2024": cplt.cplt_2024.values,
        "CPLT 2025": cplt.cplt_2025.values,
        "CPLT 2026": cplt.cplt_2026.values,
        **idx,
    }
    hojas["Descriptivos"] = pd.DataFrame(
        [
            {
                "Medición": k,
                "Media": v.mean(),
                "Mediana": np.median(v),
                "DE": v.std(ddof=1),
                "Mín.": v.min(),
                "Máx.": v.max(),
                "CV (%)": v.std(ddof=1) / v.mean() * 100,
                "GORE ≥ 90 %": int((v >= 90).sum()),
            }
            for k, v in series.items()
        ]
    )

    # Evolución CPLT
    fr = stats.friedmanchisquare(cplt.cplt_2024, cplt.cplt_2025, cplt.cplt_2026)
    W = fr.statistic / (n * 2)
    filas = [
        {
            "Prueba": "Friedman (2024, 2025, 2026)",
            "Estadístico": fr.statistic,
            "gl": 2,
            "p": fr.pvalue,
            "Tamaño del efecto": W,
            "Medida del efecto": "W de Kendall",
            "Mejoran": "",
            "Empeoran": "",
            "Δ media (p.p.)": "",
        }
    ]
    for a, b in [("cplt_2024", "cplt_2025"), ("cplt_2025", "cplt_2026"), ("cplt_2024", "cplt_2026")]:
        w = stats.wilcoxon(cplt[b], cplt[a])
        dif = cplt[b] - cplt[a]
        filas.append(
            {
                "Prueba": f"Wilcoxon {a[-4:]} → {b[-4:]}",
                "Estadístico": w.statistic,
                "gl": "",
                "p": w.pvalue,
                "Tamaño del efecto": r_efecto(w.pvalue, n),
                "Medida del efecto": "r = Z/√n",
                "Mejoran": int((dif > 0).sum()),
                "Empeoran": int((dif < 0).sum()),
                "Δ media (p.p.)": dif.mean(),
            }
        )
    for a, b in [("cplt_2024", "cplt_2025"), ("cplt_2025", "cplt_2026"), ("cplt_2024", "cplt_2026")]:
        s = stats.spearmanr(cplt[a], cplt[b])
        filas.append(
            {
                "Prueba": f"Spearman rankings {a[-4:]} – {b[-4:]}",
                "Estadístico": s.statistic,
                "gl": "",
                "p": s.pvalue,
                "Tamaño del efecto": "",
                "Medida del efecto": "ρ",
                "Mejoran": "",
                "Empeoran": "",
                "Δ media (p.p.)": "",
            }
        )
    hojas["CPLT_evolucion"] = pd.DataFrame(filas)
    w2526 = filas[2]
    resumen += [
        (
            "CPLT promedio 2024 / 2025 / 2026 (%)",
            f"{cplt.cplt_2024.mean():.1f} / {cplt.cplt_2025.mean():.1f} / {cplt.cplt_2026.mean():.1f}",
        ),
        ("Friedman CPLT", f"χ²(2) = {fr.statistic:.1f}; p = {fr.pvalue:.3f}; W = {W:.2f}"),
        (
            "Wilcoxon CPLT 2025 → 2026",
            f"p = {w2526['p']:.3f}; r = {w2526['Tamaño del efecto']:.2f}; mejoran {w2526['Mejoran']} de 16",
        ),
        ("Spearman rankings CPLT 2024 – 2026", f"ρ = {filas[6]['Estadístico']:.2f}"),
    ]

    # Convergencia CPLT 2026 vs cada scraping
    x = cplt.cplt_2026.values
    filas = []
    for k, y in idx.items():
        w = stats.wilcoxon(y, x)
        s = stats.spearmanr(x, y)
        lo, hi = spearman_boot(x, y)
        dif = y - x
        sesgo, de = dif.mean(), dif.std(ddof=1)
        po, k90 = kappa(x >= 90, y >= 90)
        filas.append(
            {
                "Scraping": k,
                "Wilcoxon p": w.pvalue,
                "Spearman ρ": s.statistic,
                "Spearman p": s.pvalue,
                "IC 95 % bootstrap inf.": lo,
                "IC 95 % bootstrap sup.": hi,
                "Bland-Altman sesgo (p.p.)": sesgo,
                "Límite inferior (p.p.)": sesgo - 1.96 * de,
                "Límite superior (p.p.)": sesgo + 1.96 * de,
                "Acuerdo ≥ 90 %": po,
                "κ ≥ 90 %": k90,
            }
        )
    hojas["Convergencia_CPLT"] = pd.DataFrame(filas)
    c1 = filas[0]
    resumen += [
        (
            "Scraping 1 vs CPLT 2026",
            f"ρ = {c1['Spearman ρ']:.2f} (IC {c1['IC 95 % bootstrap inf.']:.2f} a {c1['IC 95 % bootstrap sup.']:.2f}); "
            f"Wilcoxon p = {c1['Wilcoxon p']:.3f}; κ = {c1['κ ≥ 90 %']:.2f}",
        ),
        (
            "Bland-Altman scraping 1",
            f"sesgo {c1['Bland-Altman sesgo (p.p.)']:.1f} p.p.; límites {c1['Límite inferior (p.p.)']:.1f} a {c1['Límite superior (p.p.)']:.1f}",
        ),
    ]
    c3 = filas[2]
    resumen.append(
        (
            "Scraping 3 vs CPLT 2026",
            f"ρ = {c3['Spearman ρ']:.2f} (IC {c3['IC 95 % bootstrap inf.']:.2f} a {c3['IC 95 % bootstrap sup.']:.2f})",
        )
    )

    # Sensibilidad
    filas = []
    for k, m in zip(idx, (s1, s2, s3)):
        filas.append(
            {
                "Scraping": k,
                **{r: indice(m, r).mean() for r in ("base", "EX = 0", "EX = 0,5", "sin ER")},
                "DE (base)": indice(m).std(ddof=1),
            }
        )
    hojas["Sensibilidad"] = pd.DataFrame(filas)
    for f in filas:
        resumen.append(
            (
                f"Índice medio {f['Scraping']}",
                f"{f['base']:.1f} % (DE {f['DE (base)']:.1f}); EX = 0: {f['EX = 0']:.1f} %",
            )
        )

    # Categorías
    filas = []
    for c in CATS:
        fila = {"Cat.": c, "Categoría": NOMBRES[c]}
        for k, m in zip(("S1", "S2", "S3"), (s1, s2, s3)):
            vc = m[c].value_counts()
            for e in ("OK", "EX", "SD", "ER"):
                fila[f"{k} {e}"] = int(vc.get(e, 0))
            fila[f"{k} % disponible"] = (vc.get("OK", 0) + vc.get("EX", 0)) / 16 * 100
        filas.append(fila)
    hojas["Categorias"] = pd.DataFrame(filas)

    # Concordancia entre scrapings
    filas = []
    for (na, ma), (nb, mb) in [
        (("Scraping 1", s1), ("Scraping 2", s2)),
        (("Scraping 2", s2), ("Scraping 3", s3)),
        (("Scraping 1", s1), ("Scraping 3", s3)),
    ]:
        ea, eb = ma[CATS].values.ravel(), mb[CATS].values.ravel()
        ba, bb = np.isin(ea, ["OK", "EX"]).astype(int), np.isin(eb, ["OK", "EX"]).astype(int)
        po, k = kappa(ba, bb)
        ia, ib = indice(ma).reindex(g).values, indice(mb).reindex(g).values
        filas.append(
            {
                "Par": f"{na} – {nb}",
                "Estado idéntico (de 240)": int((ea == eb).sum()),
                "Valor 0/1 idéntico (de 240)": int((ba == bb).sum()),
                "Acuerdo": po,
                "κ": k,
                "Spearman ρ índices": stats.spearmanr(ia, ib).statistic,
                "Wilcoxon p índices": stats.wilcoxon(ib, ia).pvalue if np.any(ia != ib) else np.nan,
            }
        )
    hojas["Concordancia_scrapings"] = pd.DataFrame(filas)
    f23 = filas[1]
    resumen.append(
        (
            "Scraping 3 (código) vs scraping 2",
            f"{f23['Valor 0/1 idéntico (de 240)']} de 240 celdas; κ = {f23['κ']:.2f}; ρ = {f23['Spearman ρ índices']:.2f}",
        )
    )

    # ODS: proporciones con IC de Wilson
    web = men[men.fuente == "Sitio web"]
    cp = docs[docs.fuente == "Cuenta pública"]
    cp_disp = cp[cp.clasificacion != "No disponible"]
    ods16 = int((mapa["ODS 16"] == 2).sum())
    props = [
        ("GORE con mención ODS en su sitio web", web.gore.nunique(), 16),
        ("Menciones web retóricas", int((web.consenso == "Retórica").sum()), len(web)),
        ("ERD que mencionan los ODS (peldaño ≥ 1)", int((mapa.peldano >= 1).sum()), int(mapa.peldano.notna().sum())),
        ("ERD con vínculo sustantivo (peldaño ≥ 2)", int((mapa.peldano >= 2).sum()), int(mapa.peldano.notna().sum())),
        ("ERD operacionalizadas (peldaño 3)", int((mapa.peldano == 3).sum()), int(mapa.peldano.notna().sum())),
        ("Cuentas públicas con vínculo sustantivo", int((cp_disp.clasificacion == "Sustantiva").sum()), len(cp_disp)),
        (
            "Cuentas públicas con alguna mención",
            int(cp_disp.clasificacion.isin(["Retórica", "Sustantiva"]).sum()),
            len(cp_disp),
        ),
        ("ERD que vinculan el ODS 16", ods16, 16),
    ]
    filas = []
    for nombre, k, nn in props:
        p, lo, hi = wilson(k, nn)
        filas.append(
            {
                "Indicador": nombre,
                "k": k,
                "n": nn,
                "Proporción": p,
                "IC 95 % inf. (Wilson)": lo,
                "IC 95 % sup. (Wilson)": hi,
            }
        )
        resumen.append((nombre, f"{k} de {nn} ({p * 100:.0f} %; IC 95 %: {lo * 100:.0f}–{hi * 100:.0f} %)"))
    hojas["ODS_proporciones"] = pd.DataFrame(filas)
    hojas["ODS_mapa"] = mapa

    # ODS vs transparencia
    filas = []
    pel = mapa.set_index("gore").peldano.reindex(g).values
    ok = ~np.isnan(pel)
    for nombre, v in [("CPLT 2026", x), *idx.items()]:
        s = stats.spearmanr(pel[ok], v[ok])
        filas.append({"Peldaño ERD vs": nombre, "ρ": s.statistic, "p": s.pvalue, "n": int(ok.sum())})
    hojas["ODS_vs_transparencia"] = pd.DataFrame(filas)
    resumen.append(("Peldaño ERD vs CPLT 2026", f"ρ = {filas[0]['ρ']:.2f}; p = {filas[0]['p']:.2f}"))

    # Fiabilidad
    filas = []
    va = ver[ver.fuente == "Sitio web"].copy()
    va["coincide"] = (va.cod1 - va.cod2).abs() <= 1
    filas.append(
        {
            "Medida": "Conteos web coincidentes (±1), muestra de 5 GORE",
            "Valor": f"{int(va.coincide.sum())} de {len(va)}",
        }
    )
    vb = ver[(ver.fuente != "Sitio web") & ver.cod2.notna()].copy()
    if len(vb):
        vb["coincide"] = ((vb.cod1 - vb.cod2).abs() <= 1) | ((vb.cod1 - vb.cod2).abs() <= 0.1 * vb.cod1)
        filas.append(
            {
                "Medida": "Conteos en documentos coincidentes (±1 o ±10 %)",
                "Valor": f"{int(vb.coincide.sum())} de {len(vb)}",
            }
        )
    else:
        filas.append(
            {
                "Medida": "Conteos en documentos coincidentes (±1 o ±10 %)",
                "Valor": "segunda codificación no registrada en la planilla",
            }
        )
    par = men[men.cod2.isin(["Retórica", "Sustantiva"]) & men.cod1.isin(["Retórica", "Sustantiva"])]
    po, k = kappa(par.cod1, par.cod2)
    filas.append({"Medida": "Tipo de mención: acuerdo observado", "Valor": f"{po * 100:.1f} % (n = {len(par)})"})
    filas.append({"Medida": "Tipo de mención: κ de Cohen", "Valor": f"{k:.2f}"})
    hojas["Fiabilidad"] = pd.DataFrame(filas)
    resumen += [("Conteos web coincidentes", filas[0]["Valor"]), ("κ tipo de mención", f"{k:.2f} (n = {len(par)})")]

    # Replicación con el código
    filas = []
    if web_cod is not None:
        u_eq = {norm_url(u) for u in web.url.dropna()}
        u_cod = {norm_url(u) for u in web_cod.url.dropna()}
        comunes = u_eq & u_cod
        filas += [
            {"Medida": "Menciones web del equipo (analizadas en el paper)", "Valor": len(u_eq)},
            {"Medida": "Páginas web con mención halladas por el código", "Valor": len(u_cod)},
            {"Medida": "Coinciden (misma URL)", "Valor": len(comunes)},
            {"Medida": "Del equipo que el código no halló", "Valor": len(u_eq - u_cod)},
            {"Medida": "Halladas solo por el código (pendientes de codificar)", "Valor": len(u_cod - u_eq)},
            {
                "Medida": "GORE sin coincidencias del código",
                "Valor": ", ".join(sorted(set(web.gore) - set(web_cod.gore))),
            },
        ]
        resumen.append(
            (
                "Código vs equipo: menciones web",
                f"{len(comunes)} de {len(u_eq)} recuperadas; {len(u_cod - u_eq)} páginas nuevas",
            )
        )
    if doc_cod is not None:
        d = doc_cod[~doc_cod.estado.isin(["No disponible", "Solo video"])].copy()
        d["con_cod"] = d.estado.str.startswith("Con mención")
        ref = docs.set_index(["gore", "fuente"]).clasificacion
        d["con_eq"] = [ref.get((a, b), "") in ("Retórica", "Sustantiva") for a, b in zip(d.gore, d.fuente)]
        coinc = int((d.con_cod == d.con_eq).sum())
        filas += [
            {"Medida": "Documentos leídos por el código", "Valor": f"{len(d)} de {len(doc_cod)}"},
            {"Medida": "Presencia o ausencia de menciones igual a la del equipo", "Valor": f"{coinc} de {len(d)}"},
        ]
        resumen.append(("Código vs equipo: documentos", f"{len(d)} leídos; coincidencia {coinc} de {len(d)}"))
    hojas["Replicacion_codigo"] = pd.DataFrame(filas)

    # Potencia
    r_min = np.tanh((stats.norm.isf(0.025) + stats.norm.isf(0.20)) / np.sqrt(n - 3))
    hojas["Potencia"] = pd.DataFrame(
        [{"Medida": "Correlación mínima detectable (α = 0,05; potencia 80 %; n = 16)", "Valor": round(r_min, 2)}]
    )
    resumen.append(("Correlación mínima detectable (n = 16)", f"|ρ| ≈ {r_min:.2f}"))
    return hojas, pd.DataFrame(resumen, columns=["Indicador", "Valor"])


# ---------------------------------------------------------------- Excel
def escribir(hojas, resumen, ruta, fuentes):
    with pd.ExcelWriter(ruta, engine="openpyxl") as xw:
        resumen.to_excel(xw, sheet_name="Resumen", index=False, startrow=3)
        for nombre, df in hojas.items():
            df.to_excel(xw, sheet_name=nombre[:31], index=False)
        pd.DataFrame(fuentes, columns=["Dato", "Archivo"]).to_excel(xw, sheet_name="Fuentes", index=False)
    wb = load_workbook(ruta)
    ws = wb["Resumen"]
    ws["A1"] = "Análisis estadístico del paper — generado por analisis_paper_2026.py"
    ws["A2"] = (
        f"Fecha: {date.today():%d-%m-%Y}. Cada hoja contiene el detalle de un análisis; la hoja Fuentes indica los archivos usados."
    )
    ws["A1"].font = Font(name="Arial", size=12, bold=True)
    ws["A2"].font = Font(name="Arial", size=9, italic=True)
    azul = PatternFill("solid", fgColor="DDEBF7")
    for w in wb.worksheets:
        fila_enc = 4 if w.title == "Resumen" else 1
        for c in w[fila_enc]:
            c.font = Font(name="Arial", size=10, bold=True)
            c.fill = azul
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for fila in w.iter_rows(min_row=fila_enc + 1):
            for c in fila:
                if c.font is None or not c.font.b:
                    c.font = Font(name="Arial", size=10)
                if isinstance(c.value, float):
                    c.number_format = "0.000" if abs(c.value) < 1 else "0.0"
        for i, col in enumerate(w.columns, 1):
            largo = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            w.column_dimensions[get_column_letter(i)].width = min(max(10, largo + 2), 70)
        w.freeze_panes = w.cell(fila_enc + 1, 1)
    wb.save(ruta)


def main():
    global ENTRADA
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--planilla", default=PLANILLA)
    ap.add_argument("--entrada", default=ENTRADA, help="carpeta con portal_matriz_*.csv y las salidas ODS del código")
    a = ap.parse_args()
    ENTRADA = os.path.abspath(a.entrada)
    os.makedirs(SALIDAS, exist_ok=True)
    cplt, s1, s2, s3, et3, wb = cargar(a.planilla)
    men, docs, mapa, ver = leer_planilla(wb)
    # Web: se unen todas las corridas (buscador y mapa del sitio); documentos: la corrida más reciente
    archivos_web = sorted(glob.glob(os.path.join(ENTRADA, "ods_menciones_web_*.csv")))
    web_cod = (
        pd.concat([pd.read_csv(f, encoding="utf-8-sig") for f in archivos_web]).drop_duplicates("url")
        if archivos_web
        else None
    )
    fw = " + ".join(os.path.basename(f) for f in archivos_web)
    docs_arch = sorted(glob.glob(os.path.join(ENTRADA, "ods_documentos_*.csv")))
    fd = docs_arch[-1] if docs_arch else None
    doc_cod = pd.read_csv(fd, encoding="utf-8-sig") if fd else None
    hojas, resumen = analizar(cplt, s1, s2, s3, et3, men, docs, mapa, ver, web_cod, doc_cod)
    ruta = os.path.join(SALIDAS, f"analisis_estadistico_{date.today():%Y-%m-%d}.xlsx")
    fuentes = [
        ("CPLT 2024–2026", "referencia_cplt.csv"),
        ("Scraping 1", "referencia_scraping_2026-07-22.csv"),
        ("Scraping 2", "referencia_scraping_2026-09-25.csv"),
        ("Scraping 3", os.path.basename(ultimo("portal_matriz_*.csv"))),
        ("Codificación ODS del equipo", os.path.basename(a.planilla)),
        ("ODS web (código)", fw or "—"),
        ("ODS documentos (código)", os.path.basename(fd) if fd else "—"),
    ]
    escribir(hojas, resumen, ruta, fuentes)
    print(resumen.to_string(index=False))
    print(f"\nListo: {ruta}")


if __name__ == "__main__":
    main()
